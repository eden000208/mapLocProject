"""显式生成可重复的模拟数据，并在远程 MySQL 上验收核心链路。

仅写入 bsq_app。DEMO 标识为合成数据，不是真实用户号码。
"""

import argparse
import csv
import io
import json
import secrets
import threading
import time
import uuid
from pathlib import Path
from unittest.mock import patch

from sqlalchemy import delete, insert, select

from .services import query
from .services.auth import hash_password, verify_password
from .config import settings
from .db import SessionLocal, engine
from .detail_schema import create_detail_tables, local_table, other_table
from .main import app
from .models import (
    ImportBatch,
    LoginSession,
    OperationLog,
    QueryRow,
    QueryTask,
    Station,
    User,
)

DISTRICT = "模拟区县"
LAC = 990000
MONTHS = ["202608", "202609"]
FIELDS = ["lac", "cell_id", "station_name", "district", "service_center", "lng", "lat"]


def demo_station(station: Station) -> bool:
    """仅识别生成脚本保留的四个模拟小区，不按模糊名称删除。"""
    return (
        station.lac == LAC
        and station.cell_id in range(1, 5)
        and station.station_name == f"模拟基站{station.cell_id}"
        and station.district in {DISTRICT, "模拟未授权区县"}
        and station.service_center == ""
    )


def cleanup_demo() -> dict[str, int]:
    """清理显式模拟数据及关联任务；保留业务数据和管理员。"""
    if settings.app_database != "bsq_app" or engine.url.database != "bsq_app":
        raise ValueError("模拟数据清理仅允许在 bsq_app 执行")
    files = [
        settings.data_dir / "demo-login.json",
        settings.data_dir / "acceptance.json",
    ]
    counts = {}
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == "demo_analyst"))
        if user and (user.name != "模拟分析员" or user.role != "analyst"):
            raise ValueError("模拟账号标记已变更，停止清理")
        stations = list(db.scalars(select(Station).where(Station.lac == LAC)))
        if any(not demo_station(s) for s in stations):
            raise ValueError("模拟小区保留编号存在非模拟数据，停止清理")
        for table in (local_table, other_table):
            if (
                db.scalar(
                    select(table.c.id)
                    .where(table.c.lac == LAC, table.c.is_demo.is_(False))
                    .limit(1)
                )
                is not None
            ):
                raise ValueError("模拟小区保留编号存在业务明细，停止清理")
        tasks = list(db.scalars(select(QueryTask)))
        selected = []
        for task in tasks:
            keys = task.cell_keys or []
            has_demo = any(k.get("lac") == LAC for k in keys)
            only_demo = bool(keys) and all(
                k.get("lac") == LAC and k.get("cell_id") in range(1, 5) for k in keys
            )
            if has_demo and not only_demo:
                raise ValueError("存在混合业务与模拟小区的任务，停止清理")
            if user and task.user_id == user.id and not only_demo:
                raise ValueError("演示账号存在非模拟查询任务，停止清理")
            if only_demo or (user and task.user_id == user.id):
                if task.status in {"PENDING", "RUNNING"}:
                    raise ValueError("模拟任务仍在运行，请停止后端后清理")
                selected.append(task)
        task_ids = [t.id for t in selected]
        for task in selected:
            folder = settings.data_dir / "exports"
            files.extend(folder.glob(f"{task.id}-*-masked.xlsx"))
            files.extend(folder.glob(f"{task.id}-*-plain.xlsx"))
        if task_ids:
            db.execute(delete(QueryRow).where(QueryRow.task_id.in_(task_ids)))
            db.execute(delete(QueryTask).where(QueryTask.id.in_(task_ids)))
        if user:
            for batch in db.scalars(
                select(ImportBatch).where(ImportBatch.user_id == user.id)
            ):
                files.append(Path(batch.path))
            db.execute(delete(ImportBatch).where(ImportBatch.user_id == user.id))
            db.execute(delete(LoginSession).where(LoginSession.user_id == user.id))
            db.execute(delete(OperationLog).where(OperationLog.user_id == user.id))
            db.delete(user)
        for station in stations:
            db.delete(station)
        for table in (local_table, other_table):
            result = db.execute(delete(table).where(table.c.is_demo.is_(True)))
            counts[table.name] = result.rowcount
        root = settings.data_dir.resolve()
        if any(not p.resolve().is_relative_to(root) for p in files):
            raise ValueError("模拟文件路径超出数据目录，停止清理")
        db.commit()
    for path in files:
        path.unlink(missing_ok=True)
    counts.update(
        tasks=len(selected), stations=len(stations), users=int(user is not None)
    )
    return counts


def seed_details() -> None:
    create_detail_tables()
    with SessionLocal() as db:
        stations = db.scalars(select(Station).where(Station.lac == LAC))
        if any(not demo_station(s) for s in stations):
            raise ValueError("模拟小区保留编号存在非模拟数据，停止生成")
    with engine.begin() as connection:
        for net, table in (("local", local_table), ("other", other_table)):
            if (
                connection.scalar(
                    select(table.c.id)
                    .where(table.c.lac == LAC, table.c.is_demo.is_(False))
                    .limit(1)
                )
                is not None
            ):
                raise ValueError("模拟小区保留编号存在业务明细，停止生成")
            connection.execute(delete(table).where(table.c.is_demo.is_(True)))
            values = []
            for month in MONTHS:
                for cell in range(1, 5):
                    for index in range(30):
                        row = {
                            "month_id": month,
                            "device_number": f"DEMO-{net}-{month}-{index % 10:04d}",
                            "cell_name": f"模拟基站{cell}",
                            "street_name": "模拟街道",
                            "lac": LAC,
                            "cell_id": cell,
                            "is_demo": True,
                        }
                        if net == "local":
                            row["uid"] = f"DEMO-UID-{index % 10:04d}"
                        values.append(row)
            connection.execute(insert(table), values)


def credentials() -> dict:
    path = settings.data_dir / "demo-login.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
    else:
        data = {"username": "demo_analyst", "password": secrets.token_urlsafe(18)}
    with SessionLocal() as db:
        user = db.scalar(select(User).where(User.username == data["username"]))
        if not user:
            user = User(
                username=data["username"],
                name="模拟分析员",
                role="analyst",
                password_hash=hash_password(data["password"]),
                districts=[DISTRICT],
                service_centers=[],
                view_plain=True,
                must_change_pwd=True,
            )
            db.add(user)
            db.commit()
        elif user.name != "模拟分析员":
            raise RuntimeError("Demo username belongs to a different account")
        elif not verify_password(data["password"], user.password_hash):
            # 演示账号若已被改密，显式重新验收时生成新凭据并撤销旧会话。
            data["password"] = secrets.token_urlsafe(18)
            user.password_hash = hash_password(data["password"])
            user.failed_count = 0
            user.locked_until = None
            db.execute(delete(LoginSession).where(LoginSession.user_id == user.id))
            db.commit()
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def checked(response, status=200):
    if response.status_code != status:
        raise RuntimeError(f"Unexpected HTTP {response.status_code}, expected {status}")
    return response


def wait_task(client, task_id, expected="SUCCESS") -> dict:
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        task = checked(client.get(f"/api/queries/{task_id}")).json()
        if task["status"] in query.TERMINAL:
            if task["status"] != expected:
                raise RuntimeError(f"Task status {task['status']}, expected {expected}")
            return task
        time.sleep(0.1)
    raise RuntimeError("Acceptance task did not finish within 30 seconds")


def run_acceptance(data: dict) -> dict:
    from fastapi.testclient import TestClient

    report = {}
    temp_name = "verify_" + uuid.uuid4().hex[:12]
    temp_password = secrets.token_urlsafe(18)
    with SessionLocal() as db:
        admin = User(
            username=temp_name,
            name="临时验收管理员",
            role="admin",
            password_hash=hash_password(temp_password),
            must_change_pwd=False,
        )
        db.add(admin)
        db.commit()
        admin_id = admin.id
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            checked(client.get("/api/health"))
            login = checked(
                client.post(
                    "/api/auth/login",
                    json={
                        "username": temp_name,
                        "password": temp_password,
                    },
                )
            ).json()
            client.headers["X-CSRF-Token"] = login["csrf_token"]
            stream = io.StringIO()
            writer = csv.writer(stream)
            writer.writerow(FIELDS)
            for cell in range(1, 5):
                writer.writerow(
                    [
                        LAC,
                        cell,
                        f"模拟基站{cell}",
                        DISTRICT if cell <= 3 else "模拟未授权区县",
                        "",
                        112.20 + cell * 0.002,
                        30.98 + cell * 0.002,
                    ]
                )
            writer.writerow([LAC, 1, "模拟基站1", DISTRICT, "", 112.202, 30.982])
            writer.writerow(["invalid", 9, "模拟无效行", DISTRICT, "", 112.2, 30.98])
            uploaded = checked(
                client.post(
                    "/api/admin/stations/import",
                    files={
                        "file": (
                            "DEMO-stations.csv",
                            stream.getvalue().encode(),
                            "text/csv",
                        ),
                    },
                )
            ).json()
            batch_id = uploaded["batch_id"]
            preview = checked(
                client.post(
                    f"/api/admin/stations/import/{batch_id}/preview",
                    json={
                        "mapping": {key: key for key in FIELDS},
                        "coord_sys": "WGS84",
                    },
                )
            ).json()
            assert preview["valid_count"] == 4 and preview["error_count"] == 1
            assert preview["warning_count"] >= 1
            confirmed = checked(
                client.post(
                    f"/api/admin/stations/import/{batch_id}/confirm",
                )
            ).json()
            assert confirmed["imported"] == 4
            assert (
                checked(
                    client.post(
                        f"/api/admin/stations/import/{batch_id}/confirm",
                    )
                ).json()
                == confirmed
            )
            checked(client.get(f"/api/admin/stations/import/{batch_id}/errors"))
            report["import"] = {"valid": 4, "invalid": 1, "idempotent": True}
            admin_query = checked(
                client.post(
                    "/api/queries",
                    json={
                        "cells": [
                            {"lac": LAC, "cell_id": cell} for cell in range(1, 5)
                        ],
                        "months": MONTHS,
                        "nets": ["local", "other"],
                        "purpose": "管理员全部区域查询验收",
                    },
                )
            ).json()["id"]
            admin_task = wait_task(client, admin_query)
            assert admin_task["rows_local"] == admin_task["rows_other"] == 240
            checked(client.get(f"/api/queries/{admin_query}/rows?plain=true"))
            checked(client.post(f"/api/queries/{admin_query}/export?plain=true"))
            report["admin"] = {
                "all_active_cells": 4,
                "local_rows": 240,
                "other_rows": 240,
                "plain_export": True,
            }
            login = checked(client.post("/api/auth/login", json=data)).json()
            client.headers["X-CSRF-Token"] = login["csrf_token"]
            if login["must_change_pwd"]:
                new_password = secrets.token_urlsafe(18)
                checked(
                    client.post(
                        "/api/auth/change-password",
                        json={
                            "old_password": data["password"],
                            "new_password": new_password,
                        },
                    )
                )
                data["password"] = new_password
                (settings.data_dir / "demo-login.json").write_text(
                    json.dumps(data, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
            config = checked(client.get("/api/config")).json()
            assert config["tianditu_key"]
            months = checked(client.get("/api/meta/months?nets=local,other")).json()
            assert set(MONTHS).issubset(months["months"])
            stations = checked(
                client.get(
                    "/api/stations?bbox=112.19,30.97,112.22,31.00&zoom=17&q=模拟",
                )
            ).json()
            assert stations["items"]
            selection = checked(
                client.post(
                    "/api/selections/resolve",
                    json={
                        "type": "circle",
                        "center": [112.20, 30.98],
                        "radius": 3000,
                    },
                )
            ).json()
            assert selection["cell_count"] == 3
            keys = [{"lac": LAC, "cell_id": cell} for cell in range(1, 4)]
            checked(
                client.post(
                    "/api/selections/resolve",
                    json={
                        "type": "cells",
                        "cells": [{"lac": LAC, "cell_id": 4}],
                    },
                ),
                403,
            )
            body = {
                "cells": keys,
                "months": MONTHS,
                "nets": ["local", "other"],
                "purpose": "模拟数据核心闭环验收",
            }
            task_id = checked(client.post("/api/queries", json=body)).json()["id"]
            task = wait_task(client, task_id)
            assert task["rows_local"] == task["rows_other"] == 180
            assert task["uniq_local"] == task["uniq_other"] == 20
            masked = checked(client.get(f"/api/queries/{task_id}/rows?size=50")).json()
            assert len(masked["items"]) == 50 and masked["total"] == 180
            assert all(
                "uid" not in row and "*" in row["device_number"]
                for row in masked["items"]
            )
            assert (
                len(
                    checked(
                        client.get(
                            f"/api/queries/{task_id}/rows?page=4&size=50",
                        )
                    ).json()["items"]
                )
                == 30
            )
            checked(client.get(f"/api/queries/{task_id}/rows?plain=true"))
            exported = checked(client.post(f"/api/queries/{task_id}/export"))
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
            assert workbook.sheetnames == ["本网明细", "异网明细", "摘要"]
            workbook.close()
            checked(client.post(f"/api/queries/{task_id}/export?plain=true"))
            history = checked(client.get("/api/queries")).json()
            assert any(item["id"] == task_id for item in history["items"])
            report["query"] = {
                "task_id": task_id,
                "local_rows": 180,
                "other_rows": 180,
                "unique_per_net": 20,
                "xlsx_sheets": 3,
                "pagination": True,
            }
            with SessionLocal() as db:
                other = db.get(User, admin_id)
                other.role = "analyst"
                other.districts = [DISTRICT]
                db.commit()
            other_client = TestClient(app, raise_server_exceptions=False)
            try:
                other_login = checked(
                    other_client.post(
                        "/api/auth/login",
                        json={
                            "username": temp_name,
                            "password": temp_password,
                        },
                    )
                ).json()
                other_client.headers["X-CSRF-Token"] = other_login["csrf_token"]
                checked(other_client.get(f"/api/queries/{task_id}"), 404)
                checked(other_client.get(f"/api/queries/{task_id}/rows"), 404)
                checked(other_client.post(f"/api/queries/{task_id}/export"), 404)
                checked(other_client.post(f"/api/queries/{task_id}/cancel"), 404)
            finally:
                other_client.close()
            entered, release = threading.Event(), threading.Event()
            original_stream = query.stream_batch

            def paused_stream(*args):
                entered.set()
                if not release.wait(10):
                    raise RuntimeError("Acceptance stream was not released")
                yield from original_stream(*args)

            with patch.object(query, "stream_batch", paused_stream):
                try:
                    cancelled = checked(client.post("/api/queries", json=body)).json()[
                        "id"
                    ]
                    if not entered.wait(10):
                        raise RuntimeError("Cancellation task did not start")
                    checked(client.post(f"/api/queries/{cancelled}/cancel"))
                finally:
                    release.set()
                wait_task(client, cancelled, "CANCELLED")
            checked(client.get(f"/api/queries/{cancelled}/rows"), 409)
            user_id = login["id"]
            with SessionLocal() as db:
                user = db.get(User, user_id)
                user.view_plain = False
                db.commit()
            checked(client.get(f"/api/queries/{task_id}/rows?plain=true"), 403)
            checked(client.post(f"/api/queries/{task_id}/export?plain=true"), 403)
            with SessionLocal() as db:
                user = db.get(User, user_id)
                user.view_plain = True
                user.districts = []
                db.commit()
            checked(client.get(f"/api/queries/{task_id}/rows"), 403)
            checked(client.post(f"/api/queries/{task_id}/export"), 403)
            with SessionLocal() as db:
                user = db.get(User, user_id)
                user.districts = [DISTRICT]
                db.commit()
            previous = settings.max_rows
            try:
                settings.max_rows = 1
                limited = checked(client.post("/api/queries", json=body)).json()["id"]
                wait_task(client, limited, "FAILED")
                checked(client.get(f"/api/queries/{limited}/rows"), 409)
            finally:
                settings.max_rows = previous
            report["negative"] = {
                "unauthorized_cell": 403,
                "plain_denied": 403,
                "scope_revoked": 403,
                "row_limit_failed": True,
                "other_user_task": 404,
                "cancelled": True,
            }
            checked(client.post("/api/auth/logout"))
            checked(client.get("/api/auth/me"), 401)
            report["health_login_logout"] = True
    finally:
        with SessionLocal() as db:
            demo_user = db.scalar(select(User).where(User.username == data["username"]))
            if demo_user:
                demo_user.view_plain = True
                demo_user.districts = [DISTRICT]
                demo_user.must_change_pwd = True
            batches = list(
                db.scalars(
                    select(ImportBatch).where(
                        ImportBatch.user_id == admin_id,
                    )
                )
            )
            for batch in batches:
                Path(batch.path).unlink(missing_ok=True)
                db.delete(batch)
            db.execute(delete(LoginSession).where(LoginSession.user_id == admin_id))
            task_ids = select(QueryTask.id).where(QueryTask.user_id == admin_id)
            db.execute(delete(QueryRow).where(QueryRow.task_id.in_(task_ids)))
            db.execute(delete(QueryTask).where(QueryTask.user_id == admin_id))
            db.execute(delete(OperationLog).where(OperationLog.user_id == admin_id))
            db.execute(delete(User).where(User.id == admin_id))
            db.commit()
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--clean", action="store_true", help="清理模拟数据及关联文件")
    mode.add_argument(
        "--verify", action="store_true", help="Run real MySQL/API acceptance"
    )
    args = parser.parse_args()
    if args.clean:
        print(json.dumps(cleanup_demo(), ensure_ascii=False))
        return
    seed_details()
    data = credentials()
    if args.verify:
        report = run_acceptance(data)
        (settings.data_dir / "acceptance.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        print(json.dumps(report, ensure_ascii=False))
    print("Synthetic details seeded: 240 rows per net; no other databases changed.")
    print(
        "Demo credentials are in data/demo-login.json; first login requires password change."
    )


if __name__ == "__main__":
    main()
