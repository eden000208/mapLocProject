"""独立的管理员明细导入；不改变原有查询和基站导入流程。"""

import csv
import re
import uuid
from contextlib import contextmanager
from datetime import timedelta
from pathlib import Path

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import insert, select, text
from sqlalchemy.orm import Session

from ..audit import log_event
from ..config import settings
from ..datasource import load_sources
from ..db import engine
from ..detail_schema import local_table, other_table
from ..models import ImportBatch, User, utcnow

MAX_BYTES = 200 * 1024 * 1024
FIELDS = ("month_id", "device_number", "cell_name", "street_name", "lac", "cell_id")
PREFIX = "detail:"


def target_table(network: str):
    if network not in ("local", "other"):
        raise HTTPException(422, "请选择本网或异网")
    if settings.app_database != "bsq_app":
        raise HTTPException(422, "明细上传仅允许写入 bsq_app")
    table = local_table if network == "local" else other_table
    source = load_sources().get(network, {})
    required = (*FIELDS, "uid") if network == "local" else FIELDS
    if source.get("table") != f"bsq_app.{table.name}" or any(
        source.get("fields", {}).get(key) != key for key in required
    ):
        raise HTTPException(422, "当前查询来源与上传目标不一致，请联系管理员")
    return table


def csv_encoding(path: Path) -> str:
    # 先增量检测编码，再读取，避免 GB18030 的错误出现在导入中途。
    try:
        with path.open(encoding="utf-8-sig") as stream:
            while stream.read(1024 * 1024):
                pass
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "gb18030"


@contextmanager
def file_rows(path: Path):
    if path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if len(book.worksheets) != 1:
                raise ValueError("XLSX 文件必须只有一个工作表")
            yield book.active.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        with path.open(encoding=csv_encoding(path), newline="") as stream:
            yield csv.reader(stream)


def validate_row(row: list, headers: list[str], network: str) -> dict:
    if len(row) != len(headers):
        raise ValueError("列数与表头不一致")
    item = {}
    limits = {"device_number": 100, "uid": 200, "cell_name": 255, "street_name": 255}
    for key, value in zip(headers, row):
        if isinstance(value, float):
            if not value.is_integer():
                raise ValueError(f"{key} 不允许小数")
            value = int(value)
        value = "" if value is None else str(value).strip()
        if value.startswith("="):
            raise ValueError(f"{key} 不允许公式")
        if key == "month_id":
            if not re.fullmatch(r"[0-9]{4}(0[1-9]|1[0-2])", value):
                raise ValueError("month_id 必须为 YYYYMM 格式")
        elif key in ("lac", "cell_id"):
            if not re.fullmatch(r"[0-9]+", value) or int(value) > 9223372036854775807:
                raise ValueError(f"{key} 必须为非负 BIGINT 整数")
            item[key] = int(value)
            continue
        elif len(value) > limits[key]:
            raise ValueError(f"{key} 超出允许长度")
        if key == "device_number" and not value:
            raise ValueError("device_number 不能为空")
        item[key] = value or None
    item["is_demo"] = False
    return item


def iter_details(path: Path, network: str):
    required = (*FIELDS, "uid") if network == "local" else FIELDS
    try:
        with file_rows(path) as rows:
            headers = [str(x or "").strip() for x in next(rows)]
            if len(headers) != len(set(headers)) or set(headers) != set(required):
                raise HTTPException(422, "表头必须唯一且符合所选网别的样例列名")
            for number, row in enumerate(rows, 2):
                if all(value is None or str(value).strip() == "" for value in row):
                    continue
                try:
                    yield number, validate_row(list(row), headers, network), None
                except ValueError as exc:
                    yield number, None, str(exc)
    except HTTPException:
        raise
    except Exception:
        # 不传播文件内容、XML 或解析器异常，以免泄露真实业务标识。
        raise HTTPException(422, "文件无法读取，请检查格式、编码和工作表") from None


def scan_file(path: Path, network: str) -> dict:
    count, error_count, months, errors = 0, 0, set(), []
    for number, item, error in iter_details(path, network):
        count += 1
        if error:
            error_count += 1
            if len(errors) < 100:
                errors.append({"row": number, "message": error})
        else:
            months.add(item["month_id"])
    if not count:
        error_count = 1
        errors.append({"row": 0, "message": "文件没有数据行"})
    return {
        "network": network,
        "months": sorted(months),
        "rows": count,
        "error_count": error_count,
        "errors": errors,
        "valid": not error_count,
    }


def existing_months(db: Session, table, months: list[str], current=False) -> list[str]:
    existing = []
    for month in months:
        statement = select(table.c.id).where(table.c.month_id == month).limit(1)
        if current:
            statement = statement.with_for_update()
        if db.scalar(statement) is not None:
            existing.append(month)
    return existing


def own_batch(db: Session, batch_id: str, user: User, lock=False):
    statement = select(ImportBatch).where(ImportBatch.id == batch_id)
    if lock:
        statement = statement.with_for_update()
    batch = db.scalar(statement)
    if not batch or batch.user_id != user.id or not batch.filename.startswith(PREFIX):
        raise HTTPException(404, "明细上传批次不存在")
    if utcnow() - batch.created_at > timedelta(hours=24):
        Path(batch.path).unlink(missing_ok=True)
        raise HTTPException(410, "上传批次已过期，请重新上传")
    return batch


def clean_expired(db: Session) -> None:
    batches = db.scalars(
        select(ImportBatch).where(
            ImportBatch.filename.startswith(PREFIX),
            ImportBatch.created_at < utcnow() - timedelta(hours=24),
        )
    )
    for batch in batches:
        Path(batch.path).unlink(missing_ok=True)
        db.delete(batch)
    db.commit()


def upload(file: UploadFile, network: str, user: User, db: Session):
    target_table(network)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".xlsx", ".csv"):
        raise HTTPException(422, "仅支持 CSV 或 XLSX 文件")
    clean_expired(db)
    batch_id = str(uuid.uuid4())
    folder = settings.data_dir / "detail_uploads"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{batch_id}{suffix}"
    size = 0
    try:
        with path.open("xb") as stream:
            while chunk := file.file.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_BYTES:
                    raise HTTPException(413, "文件最大允许 200 MB")
                stream.write(chunk)
        if not size:
            raise HTTPException(422, "文件为空")
        batch = ImportBatch(
            id=batch_id,
            user_id=user.id,
            filename=f"{PREFIX}{network}",
            path=str(path),
            status="DETAIL_UPLOADED",
            preview={"network": network},
        )
        db.add(batch)
        db.commit()
    except Exception:
        db.rollback()
        path.unlink(missing_ok=True)
        raise
    finally:
        file.file.close()
    return {"batch_id": batch_id, "network": network}


def preview(batch_id: str, user: User, db: Session):
    batch = own_batch(db, batch_id, user, lock=True)
    if batch.status == "DETAIL_CONFIRMED":
        raise HTTPException(409, "该批次已导入")
    network = batch.preview["network"]
    table = target_table(network)
    result = scan_file(Path(batch.path), network)
    result["existing_months"] = existing_months(db, table, result["months"])
    result["valid"] = result["valid"] and not result["existing_months"]
    batch.preview = result
    batch.status = "DETAIL_PREVIEWED"
    db.commit()
    return result


def write_batch(db: Session, batch_id: str, user: User) -> dict:
    batch = own_batch(db, batch_id, user, lock=True)
    if batch.status != "DETAIL_PREVIEWED" or not batch.preview.get("valid"):
        raise HTTPException(409, "请先完成无错误的校验预览，且不可重复确认")
    network = batch.preview["network"]
    table = target_table(network)
    result = scan_file(Path(batch.path), network)
    if not result["valid"] or result["months"] != batch.preview["months"]:
        raise HTTPException(422, "文件校验失败，请重新上传")
    if existing_months(db, table, result["months"], current=True):
        raise HTTPException(409, "文件包含已有数据的月账期，整批拒绝")
    buffer = []
    for _, item, error in iter_details(Path(batch.path), network):
        if error:
            raise HTTPException(422, "文件校验失败，整批未导入")
        buffer.append(item)
        if len(buffer) == 1000:
            db.execute(insert(table), buffer)
            buffer = []
    if buffer:
        db.execute(insert(table), buffer)
    batch.status = "DETAIL_CONFIRMED"
    batch.preview = result
    log_event(
        db,
        user.id,
        "detail_import",
        {
            "network": network,
            "months": result["months"],
            "rows": result["rows"],
        },
    )
    db.commit()
    return {"imported": result["rows"], "months": result["months"]}


def confirm(batch_id: str, user: User, db: Session):
    batch = own_batch(db, batch_id, user)
    network, path = batch.preview["network"], Path(batch.path)
    target_table(network)
    db.rollback()
    # 固定 MySQL 连接持有网别锁，串行化同网别确认；成功前只提交一次。
    with engine.connect() as connection:
        lock_name = f"bsq_app.detail_import.{network}"
        acquired = connection.scalar(
            text("SELECT GET_LOCK(:name, 0)"), {"name": lock_name}
        )
        connection.commit()
        if acquired != 1:
            raise HTTPException(409, "该网别正在导入，请稍后重试")
        try:
            with Session(bind=connection) as writer:
                result = write_batch(writer, batch_id, user)
        finally:
            connection.rollback()
            connection.execute(text("SELECT RELEASE_LOCK(:name)"), {"name": lock_name})
            connection.commit()
    path.unlink(missing_ok=True)
    return result
