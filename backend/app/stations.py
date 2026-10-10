import csv
import io
import math
import threading
import uuid
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from shapely.geometry import Point, box, shape
from shapely.strtree import STRtree
from sqlalchemy import func, or_, select, tuple_
from sqlalchemy.dialects.mysql import insert
from sqlalchemy.orm import Session

from .audit import log_event
from .auth import cell_visible, require_admin, require_query
from .config import settings
from .db import get_db
from .geo import distance_m, to_wgs, valid_coordinate
from .models import ImportBatch, Station, User, utcnow

router = APIRouter()
_lock = threading.RLock()
_index = None
FIELDS = ("lac", "cell_id", "station_name", "district", "service_center", "lng", "lat")
ALIASES = {
    "lac": ("lac", "LAC"),
    "cell_id": ("cell_id", "CELL_ID", "小区ID", "小区编号", "小区号", "CI"),
    "station_name": ("station_name", "基站名称", "小区名称"),
    "district": ("district", "区县", "基站归属区县"),
    "service_center": ("service_center", "营服", "基站归属营服"),
    "lng": ("lng", "经度", "longitude", "REST_LONGITUDE"),
    "lat": ("lat", "纬度", "latitude", "REST_LATITUDE"),
}


def invalidate_index() -> None:
    global _index
    with _lock:
        _index = None


def get_index(db: Session):
    global _index
    with _lock:
        if _index is None:
            rows = list(db.scalars(select(Station).where(Station.is_active.is_(True))))
            snapshots = [cell_payload(s) for s in rows]
            points = [Point(s["lng"], s["lat"]) for s in snapshots]
            _index = (snapshots, STRtree(points))
        return _index


def cell_payload(s: Station) -> dict:
    return {k: getattr(s, k) for k in ("id", *FIELDS)} | {"is_active": s.is_active}


def _visible(user: User, row: dict) -> bool:
    if user.role == "admin":
        return True
    return row["district"] in (user.districts or []) or row["service_center"] in (
        user.service_centers or []
    )


def validate_cells(db: Session, user: User, keys: list) -> list[Station]:
    try:
        wanted = {(parse_integer(k["lac"]), parse_integer(k["cell_id"])) for k in keys}
    except (KeyError, ValueError, TypeError):
        raise HTTPException(422, "小区参数无效") from None
    if not wanted or len(wanted) > settings.max_cells:
        raise HTTPException(422, "小区数量超出限制")
    stations = []
    wanted_list = sorted(wanted)
    for start in range(0, len(wanted_list), 500):
        stations.extend(
            db.scalars(
                select(Station).where(
                    tuple_(Station.lac, Station.cell_id).in_(
                        wanted_list[start : start + 500]
                    )
                )
            )
        )
    if len(stations) != len(wanted) or any(not s.is_active for s in stations):
        raise HTTPException(422, "小区不存在或已停用")
    if any(not cell_visible(user, s) for s in stations):
        raise HTTPException(403, "所选小区超出当前授权范围")
    return stations


def read_upload(path: Path) -> tuple[list[str], list[list]]:
    if path.suffix == ".xlsx":
        try:
            with path.open("rb") as stream:
                book = load_workbook(stream, read_only=True, data_only=True)
                try:
                    iterator = book.active.iter_rows(values_only=True)
                    headers = [str(x or "").strip() for x in next(iterator)]
                    rows = []
                    for row in iterator:
                        if len(rows) >= 100000:
                            raise ValueError("导入最多100000行")
                        if any(x is not None for x in row):
                            rows.append(list(row))
                finally:
                    book.close()
        except (OSError, StopIteration, ValueError) as exc:
            raise HTTPException(422, str(exc) or "文件为空") from None
        except Exception:
            raise HTTPException(422, "无法读取XLSX文件") from None
    else:
        try:
            text = path.read_bytes().decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = path.read_bytes().decode("gb18030")
            except UnicodeDecodeError:
                raise HTTPException(422, "CSV编码必须为UTF-8或GB18030") from None
        rows_all = list(csv.reader(io.StringIO(text)))
        if not rows_all:
            raise HTTPException(422, "文件为空")
        headers = [v.strip() for v in rows_all[0]]
        rows = [r for r in rows_all[1:] if any(r)]
        if len(rows) > 100000:
            raise HTTPException(422, "导入最多100000行")
    if not headers or len(headers) > 100 or len(set(headers)) != len(headers):
        raise HTTPException(422, "表头必须唯一，且不超过100列")
    return headers, rows


def parse_integer(value) -> int:
    try:
        n = Decimal(str(value).strip())
        if (
            not n.is_finite()
            or n != n.to_integral_value()
            or not 0 <= n <= 9223372036854775807
        ):
            raise ValueError()
        return int(n)
    except (ValueError, InvalidOperation):
        raise ValueError("lac/cell_id 必须为非负整数") from None


def parse_rows(headers, rows, mapping, coord_sys):
    if any(mapping.get(k) not in headers for k in FIELDS):
        raise HTTPException(422, "请映射全部必需字段")
    cols = {key: headers.index(mapping[key]) for key in FIELDS}
    valid, errors, warnings = {}, [], []
    for row_no, row in enumerate(rows, 2):
        try:
            values = {key: row[i] if i < len(row) else "" for key, i in cols.items()}
            for key in ("lac", "cell_id"):
                values[key] = parse_integer(values[key])
            for key in ("station_name", "district", "service_center"):
                values[key] = str(values[key] or "").strip()
                if len(values[key]) > (200 if key == "station_name" else 100):
                    raise ValueError("名称超过字段长度")
            if not values["district"] and not values["service_center"]:
                warnings.append(
                    {"row": row_no, "message": "缺少区县和营服，分析员不可见"}
                )
            lng, lat = float(values["lng"]), float(values["lat"])
            values["lng"], values["lat"] = to_wgs(lng, lat, coord_sys)
            values.update(lng_raw=lng, lat_raw=lat, coord_sys=coord_sys, is_active=True)
            key = (values["lac"], values["cell_id"])
            if key in valid:
                warnings.append({"row": row_no, "message": "重复小区，采用最后一行"})
            valid[key] = values
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append(
                {
                    "row": row_no,
                    "message": str(exc)
                    if isinstance(exc, ValueError)
                    else "字段格式无效",
                }
            )
    return list(valid.values()), errors, warnings


class PreviewBody(BaseModel):
    mapping: dict[str, str]
    coord_sys: Literal["WGS84", "GCJ02", "BD09"] = "WGS84"


def own_batch(db, batch_id, user, lock=False):
    stmt = select(ImportBatch).where(
        ImportBatch.id == batch_id, ImportBatch.user_id == user.id
    )
    batch = db.scalar(stmt.with_for_update() if lock else stmt)
    if not batch:
        raise HTTPException(404, "导入不存在")
    return batch


def csv_response(rows, filename):
    stream = io.StringIO()
    writer = csv.writer(stream)
    for row in rows:
        writer.writerow(
            [
                "'" + str(v) if str(v).startswith(("=", "+", "-", "@")) else v
                for v in row
            ]
        )
    return Response(
        content="\ufeff" + stream.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/stations/template")
def template(user: User = Depends(require_admin)):
    return csv_response([FIELDS], "stations-template.csv")


@router.post("/admin/stations/import")
async def upload(
    file: UploadFile = File(...),
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise HTTPException(422, "只支持 CSV 和 XLSX")
    batch_id = str(uuid.uuid4())
    folder = settings.data_dir / "imports"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / (batch_id + suffix)
    size = 0
    try:
        with path.open("wb") as stream:
            while chunk := await file.read(65536):
                size += len(chunk)
                if size > settings.upload_max_mb * 1024 * 1024:
                    raise HTTPException(413, "文件超过大小限制")
                stream.write(chunk)
        headers, rows = read_upload(path)
        batch = ImportBatch(
            id=batch_id,
            user_id=user.id,
            filename=(file.filename or "")[:255],
            path=str(path.resolve()),
            status="UPLOADED",
        )
        db.add(batch)
        db.commit()
    except Exception:
        path.unlink(missing_ok=True)
        raise
    mapping = {
        key: next((h for h in aliases if h in headers), "")
        for key, aliases in ALIASES.items()
    }
    return {
        "batch_id": batch_id,
        "headers": headers,
        "mapping": mapping,
        "samples": [[str(x or "") for x in r] for r in rows[:5]],
    }


@router.post("/admin/stations/import/{batch_id}/preview")
def preview(
    batch_id: str,
    body: PreviewBody,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    batch = own_batch(db, batch_id, user, True)
    if batch.status == "CONFIRMED":
        raise HTTPException(409, "此导入已确认")
    headers, rows = read_upload(Path(batch.path))
    valid, errors, warnings = parse_rows(headers, rows, body.mapping, body.coord_sys)
    existing = 0
    for start in range(0, len(valid), 500):
        keys = [(r["lac"], r["cell_id"]) for r in valid[start : start + 500]]
        existing += (
            db.scalar(
                select(func.count())
                .select_from(Station)
                .where(tuple_(Station.lac, Station.cell_id).in_(keys))
            )
            or 0
        )
    points = [
        {
            "station_name": r["station_name"],
            "raw": [r["lng_raw"], r["lat_raw"]],
            **{
                coord: list(to_wgs(r["lng_raw"], r["lat_raw"], coord))
                for coord in ("WGS84", "GCJ02", "BD09")
            },
        }
        for r in valid[:50]
    ]
    report = {
        "total": len(rows),
        "valid_count": len(valid),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "new_count": len(valid) - existing,
        "update_count": existing,
        "errors": errors[:1000],
        "warnings": warnings[:1000],
        "points": points,
    }
    batch.preview = {
        "mapping": body.mapping,
        "coord_sys": body.coord_sys,
        "report": report,
    }
    batch.status = "PREVIEWED"
    db.commit()
    return report


@router.post("/admin/stations/import/{batch_id}/confirm")
def confirm(
    batch_id: str, user: User = Depends(require_admin), db: Session = Depends(get_db)
):
    batch = own_batch(db, batch_id, user, True)
    if batch.status == "CONFIRMED":
        return batch.preview["confirmed"]
    if not batch.preview or batch.status != "PREVIEWED":
        raise HTTPException(409, "请先预览")
    headers, rows = read_upload(Path(batch.path))
    valid, errors, _ = parse_rows(
        headers, rows, batch.preview["mapping"], batch.preview["coord_sys"]
    )
    if not valid:
        raise HTTPException(422, "没有有效行")
    for start in range(0, len(valid), 500):
        stmt = insert(Station).values(valid[start : start + 500])
        updates = {
            k: getattr(stmt.inserted, k)
            for k in valid[0]
            if k not in ("lac", "cell_id")
        }
        stmt = stmt.on_duplicate_key_update(**updates, updated_at=utcnow())
        db.execute(stmt)
    result = {"imported": len(valid), "error_count": len(errors)}
    batch.preview = batch.preview | {"confirmed": result}
    batch.status = "CONFIRMED"
    log_event(db, user.id, "import_stations", {"batch_id": batch.id, **result})
    db.commit()
    invalidate_index()
    return result


@router.get("/admin/stations/import/{batch_id}/errors")
def errors_csv(
    batch_id: str, user: User = Depends(require_admin), db: Session = Depends(get_db)
):
    batch = own_batch(db, batch_id, user)
    if not batch.preview:
        raise HTTPException(409, "请先预览")
    headers, rows = read_upload(Path(batch.path))
    _, errors, _ = parse_rows(
        headers, rows, batch.preview["mapping"], batch.preview["coord_sys"]
    )
    return csv_response(
        [["row", "message"]] + [[e["row"], e["message"]] for e in errors],
        "import-errors.csv",
    )


@router.get("/admin/stations")
def admin_stations(
    q: str = "",
    page: int = 1,
    size: int = 50,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    page, size = max(1, page), min(200, max(1, size))
    stmt = select(Station)
    if q:
        stmt = stmt.where(
            or_(
                Station.station_name.contains(q),
                Station.lac == int(q) if q.isdigit() else False,
                Station.cell_id == int(q) if q.isdigit() else False,
            )
        )
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.scalars(stmt.order_by(Station.id).offset((page - 1) * size).limit(size))
    return {
        "items": [
            cell_payload(s)
            | {"coord_sys": s.coord_sys, "lng_raw": s.lng_raw, "lat_raw": s.lat_raw}
            for s in rows
        ],
        "total": total,
    }


class StationPatch(BaseModel):
    station_name: str | None = Field(default=None, max_length=200)
    district: str | None = Field(default=None, max_length=100)
    service_center: str | None = Field(default=None, max_length=100)
    lng: float | None = Field(default=None, ge=-180, le=180)
    lat: float | None = Field(default=None, ge=-90, le=90)
    is_active: bool | None = None


@router.patch("/admin/stations/{station_id}")
def edit_station(
    station_id: int,
    body: StationPatch,
    user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    station = db.get(Station, station_id)
    if not station:
        raise HTTPException(404, "小区不存在")
    values = body.model_dump(exclude_unset=True, exclude_none=True)
    if not valid_coordinate(
        values.get("lng", station.lng), values.get("lat", station.lat)
    ):
        raise HTTPException(422, "经纬度无效")
    for key, value in values.items():
        setattr(station, key, value)
    if "lng" in values or "lat" in values:
        station.lng_raw, station.lat_raw, station.coord_sys = (
            station.lng,
            station.lat,
            "WGS84",
        )
    log_event(db, user.id, "edit_station", {"station_id": station_id})
    db.commit()
    invalidate_index()
    return cell_payload(station)


@router.get("/stations")
def map_stations(
    bbox: str = "-180,-90,180,90",
    zoom: int = 12,
    q: str = "",
    user: User = Depends(require_query),
    db: Session = Depends(get_db),
):
    try:
        west, south, east, north = [float(x) for x in bbox.split(",")]
        if (
            not valid_coordinate(west, south)
            or not valid_coordinate(east, north)
            or west >= east
            or south >= north
        ):
            raise ValueError()
    except ValueError:
        raise HTTPException(422, "视口坐标无效") from None
    rows, tree = get_index(db)
    selected = [rows[int(i)] for i in tree.query(box(west, south, east, north))]
    selected = [
        r
        for r in selected
        if _visible(user, r)
        and (
            not q
            or q.lower() in r["station_name"].lower()
            or q in (str(r["lac"]), str(r["cell_id"]))
        )
    ]
    sites = {}
    for row in selected:
        key = (row["station_name"], row["lng"], row["lat"])
        sites.setdefault(key, []).append(row)
    # 通过地理网格聚合控制响应大小，避免静默丢弃小区。
    step = 360 / (2 ** min(20, max(0, zoom))) * 60 / 256
    buckets = {}
    for (name, lng, lat), cells in sites.items():
        key = (math.floor(lng / step), math.floor(lat / step))
        buckets.setdefault(key, []).append((name, lng, lat, cells))
    items = []
    for bucket in buckets.values():
        if len(bucket) == 1:
            name, lng, lat, cells = bucket[0]
            items.append(
                {
                    "kind": "station",
                    "station_name": name,
                    "lng": lng,
                    "lat": lat,
                    "cells": cells,
                }
            )
        else:
            xs, ys = [s[1] for s in bucket], [s[2] for s in bucket]
            items.append(
                {
                    "kind": "cluster",
                    "lng": sum(xs) / len(xs),
                    "lat": sum(ys) / len(ys),
                    "count": len(bucket),
                    "bbox": [min(xs), min(ys), max(xs), max(ys)],
                }
            )
    return {"items": items, "truncated": False}


class SelectionBody(BaseModel):
    type: Literal["rect", "polygon", "circle", "cells"]
    geometry: dict | None = None
    center: tuple[float, float] | None = None
    radius: float | None = Field(default=None, gt=0, le=100000)
    cells: list[dict] | None = Field(default=None, max_length=100000)


@router.post("/selections/resolve")
def resolve(
    body: SelectionBody,
    user: User = Depends(require_query),
    db: Session = Depends(get_db),
):
    if body.type == "cells":
        cells = [cell_payload(s) for s in validate_cells(db, user, body.cells or [])]
        excluded = 0
    else:
        rows, tree = get_index(db)
        try:
            if body.type == "circle":
                if (
                    not body.center
                    or not body.radius
                    or not valid_coordinate(*body.center)
                ):
                    raise ValueError()
                lng, lat = body.center
                delta = body.radius / 110000
                dx = min(180, delta / max(0.01, math.cos(math.radians(lat))))
                geom = box(lng - dx, lat - delta, lng + dx, lat + delta)
            else:
                geom = shape(body.geometry or {})
                if geom.geom_type != "Polygon" or not geom.is_valid or geom.is_empty:
                    raise ValueError()
                if len(geom.exterior.coords) > 1000 or not all(
                    valid_coordinate(*p) for p in geom.exterior.coords
                ):
                    raise ValueError()
            candidates = [rows[int(i)] for i in tree.query(geom)]
            hits = [
                r
                for r in candidates
                if (
                    distance_m(body.center, (r["lng"], r["lat"])) <= body.radius
                    if body.type == "circle"
                    else geom.covers(Point(r["lng"], r["lat"]))
                )
            ]
        except (ValueError, TypeError, KeyError):
            raise HTTPException(422, "选择几何无效") from None
        cells = [r for r in hits if _visible(user, r)]
        excluded = len(hits) - len(cells)
    if len(cells) > settings.max_cells:
        raise HTTPException(422, f"最多选择{settings.max_cells}个小区，请缩小范围")
    return {
        "cells": cells,
        "cell_count": len(cells),
        "excluded_count": excluded,
        "station_count": len({(r["station_name"], r["lng"], r["lat"]) for r in cells}),
    }
